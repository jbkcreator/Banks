"""MOD-01 Slack file intake — should_ingest_file gate + CSV/PDF/docx integration.

Decisions (2026-09-01):
- ALL uploads (CSV, PDF, docx) now require @banks tag (app_mention event).
- should_ingest_mention_file: new gate for app_mention file uploads.
- PDF/docx: local text extraction → existing manual_intake LLM extractor.
- Too-little-text guard: < 200 chars → tell Josh (likely a scanned image PDF).
- One receipt per file; loop over multiple drops in one message.
"""
from __future__ import annotations

import dataclasses
import os
import tempfile

import pytest

from banks.chatport import FakeChatPort
from banks.config import load_config
from banks.csvport import FakeCSVPort
from banks.intake import ingest_simplify
from banks.socket_listener import should_ingest_file, should_ingest_mention_file
from banks.store import init_db


def _cfg():
    return dataclasses.replace(
        load_config(), approver_user_id="UJOSH", slack_jobs_channel_id="CJOBS",
    )


def _evt(**kw):
    e = {"user": "UJOSH", "channel": "CJOBS",
         "files": [{"name": "simplify.csv", "filetype": "csv"}]}
    e.update(kw)
    return e


def test_gate_accepts_josh_csv_in_jobs_channel():
    assert should_ingest_file(_cfg(), _evt()) is True


def test_gate_rejects_wrong_channel():
    assert should_ingest_file(_cfg(), _evt(channel="CRANDOM")) is False


def test_gate_rejects_unauthorized_user():
    assert should_ingest_file(_cfg(), _evt(user="USOMEONE")) is False


def test_gate_rejects_bot_file():
    assert should_ingest_file(_cfg(), _evt(bot_id="B1")) is False


def test_gate_rejects_non_csv():
    assert should_ingest_file(
        _cfg(), _evt(files=[{"name": "resume.pdf", "filetype": "pdf"}])
    ) is False


def test_gate_rejects_no_files():
    assert should_ingest_file(_cfg(), _evt(files=[])) is False


# --- intake integration: a Simplify export produces the expected counts -------

@pytest.fixture
def db_path():
    p = os.path.join(tempfile.mkdtemp(), "t.db")
    init_db(p)
    return p


def _simplify_rows():
    return [
        {"Job Title": "VP Sales", "Company Name": "Acme", "Location": "Remote",
         "Job URL": "https://x.com/1", "Applied Date": "2026-08-01", "Status": "Applied"},
        {"Job Title": "AE", "Company Name": "Beta", "Location": "Tampa, FL",
         "Job URL": "https://x.com/2", "Applied Date": "2026-08-02", "Status": "Applied"},
        # duplicate of row 1 (same URL)
        {"Job Title": "VP Sales", "Company Name": "Acme", "Location": "Remote",
         "Job URL": "https://x.com/1", "Applied Date": "2026-08-01", "Status": "Applied"},
    ]


def test_ingest_simplify_counts(db_path):
    res = ingest_simplify(db_path, FakeCSVPort(_simplify_rows()), "ignored", FakeChatPort())
    assert res.ingested == 2
    assert res.duplicates == 1
    # Simplify rows have no industry → held for enrichment, not surfaced
    assert res.held == 2
    assert res.surfaced == 0


# ---------------------------------------------------------------------------
# @banks-tagged file uploads (app_mention gate)
# All uploads now require @banks; should_ingest_mention_file is the new gate.
# ---------------------------------------------------------------------------

def _mention_evt(**kw):
    e = {"user": "UJOSH", "channel": "CJOBS",
         "files": [{"name": "simplify.csv", "filetype": "csv"}]}
    e.update(kw)
    return e


class TestMentionFileGate:
    def test_accepts_csv_with_mention(self):
        assert should_ingest_mention_file(_cfg(), _mention_evt()) is True

    def test_accepts_pdf_with_mention(self):
        e = _mention_evt(files=[{"name": "jd.pdf", "filetype": "pdf"}])
        assert should_ingest_mention_file(_cfg(), e) is True

    def test_accepts_docx_with_mention(self):
        e = _mention_evt(files=[{"name": "jd.docx", "filetype": "docx"}])
        assert should_ingest_mention_file(_cfg(), e) is True

    def test_rejects_unauthorized_user(self):
        assert should_ingest_mention_file(_cfg(), _mention_evt(user="UBAD")) is False

    def test_rejects_bot_file(self):
        assert should_ingest_mention_file(_cfg(), _mention_evt(bot_id="B1")) is False

    def test_rejects_no_files(self):
        assert should_ingest_mention_file(_cfg(), _mention_evt(files=[])) is False

    def test_old_gate_still_rejects_non_csv(self):
        # should_ingest_file (message event, legacy) never accepts PDF/docx
        e = _evt(files=[{"name": "jd.pdf", "filetype": "pdf"}])
        assert should_ingest_file(_cfg(), e) is False


# ---------------------------------------------------------------------------
# PDF / docx text extraction
# ---------------------------------------------------------------------------

class TestDocumentExtraction:
    def test_extract_pdf_text_returns_string(self):
        from banks.docparse import extract_text
        # Minimal valid one-line "PDF" that pypdf reads (use a known-good fixture).
        # For unit test we use a tiny synthetic PDF created at test time.
        try:
            import fpdf  # optional; skip if not installed
        except ImportError:
            pytest.skip("fpdf not installed — use pre-built fixture instead")
        pdf = fpdf.FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(200, 10, txt="Senior VP Sales at Acme Corp. Salary $250k.")
        data = pdf.output()
        text = extract_text(data, "application.pdf")
        assert "Acme" in text or len(text) > 10

    def test_extract_docx_text_returns_string(self):
        from banks.docparse import extract_text
        try:
            from docx import Document as _D  # python-docx
        except ImportError:
            pytest.skip("python-docx not installed")
        import io
        doc = _D()
        doc.add_paragraph("Looking for a VP Sales leader. Base $200k.")
        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()
        text = extract_text(data, "jd.docx", min_chars=0)
        assert "VP Sales" in text

    def test_too_little_text_raises(self):
        from banks.docparse import TooLittleText, extract_text
        # Fewer than 200 chars → TooLittleText
        try:
            import fpdf
        except ImportError:
            pytest.skip("fpdf not installed")
        pdf = fpdf.FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(200, 10, txt="Short.")  # way under 200 chars
        data = pdf.output()
        with pytest.raises(TooLittleText):
            extract_text(data, "scan.pdf", min_chars=200)

    def test_unknown_extension_raises(self):
        from banks.docparse import extract_text
        with pytest.raises(ValueError, match="unsupported"):
            extract_text(b"data", "report.xlsx")

    def test_extract_text_plain_txt(self):
        from banks.docparse import extract_text
        data = b"VP of Sales role at Acme. Remote. $250k base."
        text = extract_text(data, "jd.txt", min_chars=0)
        assert "VP of Sales" in text
